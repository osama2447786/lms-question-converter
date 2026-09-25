import streamlit as st
import docx
from docx.document import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph
import pandas as pd
import io
import re

# --- Helper functions to parse Word structure ---
def iter_block_items(parent):
    """Iterates through paragraphs and tables in their actual document order."""
    if isinstance(parent, Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("Invalid parent object")

    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)

def extract_colored_lines(block):
    """Extracts text line-by-line, detecting soft-returns and red text formatting."""
    lines = []
    current_line_text = ""
    current_line_red = False
    
    for run in block.runs:
        run_text = run.text
        is_red = False
        # Detect Red text (hex code FF0000)
        if run.font.color and run.font.color.rgb and str(run.font.color.rgb) == 'FF0000':
            is_red = True
            
        parts = run_text.split('\n')
        for idx, part in enumerate(parts):
            current_line_text += part
            if is_red and part.strip():
                current_line_red = True
                
            # If a newline occurs, save the collected line and reset
            if idx < len(parts) - 1: 
                if current_line_text.strip():
                    lines.append({'text': current_line_text.strip(), 'is_red': current_line_red})
                current_line_text = ""
                current_line_red = False
                
    # Catch any trailing text
    if current_line_text.strip():
        lines.append({'text': current_line_text.strip(), 'is_red': current_line_red})
        
    return lines

def is_new_question_start(text):
    """Heuristic logic to detect if a line of text is a new question."""
    text = text.strip()
    if not text: return False
    
    # 1. Explicit numbering (e.g., "1. ", "1)", "Q1.")
    if re.match(r'^(Q?\d+[\.\)])\s+', text, re.IGNORECASE): 
        return True
        
    # 2. Strong question punctuation
    if text.endswith('?') or text.endswith('=') or text.endswith('...') or text.endswith(':'): 
        return True
    if text.startswith('___'): 
        return True
        
    # 3. Common question starting words
    if re.match(r'^(which|what|how|why|when|where|name|select)\b', text, re.IGNORECASE) and len(text) > 15: 
        return True
        
    # 4. Specific Machinist/Technical test endings
    if text.endswith("is") or text.endswith("called") or text.endswith("cost"): 
        return True
        
    # 5. Catch-all: If it's a very long sentence and doesn't look like an option
    is_option_format = re.match(r'^([a-zA-Z][\.\)]\s+|Answer:)', text, re.IGNORECASE)
    if not is_option_format and len(text) > 55:
        return True
        
    return False

# --- Core Dynamic Parsing Logic ---
def parse_unlimited_options(file_obj, exam_id):
    doc = docx.Document(file_obj)
    blocks = list(iter_block_items(doc))
    
    # Flatten the document into a strict line-by-line sequence
    all_lines = []
    parsing_started = False
    
    for block in blocks:
        if isinstance(block, Paragraph):
            if "Part-1:" in block.text:
                parsing_started = True
                continue
            if not parsing_started:
                continue
                
            has_img = 'w:drawing' in block._element.xml or 'v:imagedata' in block._element.xml
            para_lines = extract_colored_lines(block)
            for l in para_lines:
                l['has_image'] = has_img
                all_lines.append(l)
                
        elif isinstance(block, Table):
            if not parsing_started: continue
            has_img = 'w:drawing' in block._element.xml or 'v:imagedata' in block._element.xml
            for row in block.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        cell_lines = extract_colored_lines(p)
                        for l in cell_lines:
                            l['has_image'] = has_img
                            all_lines.append(l)

    # Group the flattened lines into dynamic Question/Option packages
    questions = []
    current_q = None
    current_options = []
    correct_opt = None
    q_has_image = False
    
    for line in all_lines:
        text = line['text']
        is_red = line['is_red']
        has_img = line['has_image']
        
        starts_with_num = bool(re.match(r'^(Q?\d+[\.\)])\s+', text, re.IGNORECASE))
        looks_like_q = is_new_question_start(text)
        is_option_format = bool(re.match(r'^([a-zA-Z][\.\)]\s+|Answer:)', text, re.IGNORECASE))
        
        # Decide if this line breaks off into a new question
        is_new = False
        if current_q is None:
            is_new = True
        elif starts_with_num:
            is_new = True
        elif len(current_options) >= 1:
            if looks_like_q and not is_option_format:
                is_new = True
            elif is_option_format:
                is_new = False
            elif is_red:
                is_new = False # Red text is always the correct option
            elif len(current_options) >= 2 and not is_option_format and len(text) > 45:
                # Ambiguous long text after options -> assume new question
                is_new = True

        if is_new:
            # Package the PREVIOUS question before starting the new one
            if current_q:
                # Enforce a default correct answer if no red text was found
                if not correct_opt and current_options:
                    correct_opt = current_options[0] 
                    
                if q_has_image:
                    questions.append({
                        'q_id': f"{exam_id}_{len(questions) + 1}",
                        'text': f"[IMAGE REQUIRED - UPDATE MANUALLY] {current_q}",
                        'options': current_options if current_options else ["Option A", "Option B", "Option C", "Option D"],
                        'correct': correct_opt if current_options else "Option A"
                    })
                elif len(current_options) > 1:
                    questions.append({
                        'q_id': f"{exam_id}_{len(questions) + 1}",
                        'text': current_q,
                        'options': current_options,
                        'correct': correct_opt
                    })
                    
            # Reset trackers for the NEW question
            current_q = text
            current_options = []
            correct_opt = None
            q_has_image = has_img
        else:
            # Add to current options
            current_options.append(text)
            if is_red: 
                correct_opt = text
                
        if has_img:
            q_has_image = True

    # Package the FINAL question in the loop
    if current_q:
        if not correct_opt and current_options:
            correct_opt = current_options[0]
        if q_has_image:
            questions.append({
                'q_id': f"{exam_id}_{len(questions) + 1}",
                'text': f"[IMAGE REQUIRED - UPDATE MANUALLY] {current_q}",
                'options': current_options if current_options else ["Option A", "Option B", "Option C", "Option D"],
                'correct': correct_opt if current_options else "Option A"
            })
        elif len(current_options) > 1:
            questions.append({
                'q_id': f"{exam_id}_{len(questions) + 1}",
                'text': current_q,
                'options': current_options,
                'correct': correct_opt
            })
            
    return questions

# --- Excel Generation Logic ---
def generate_lms_excel(parsed_data, domain_id):
    rows = []
    for q in parsed_data:
        for opt_idx, option_text in enumerate(q['options']):
            is_correct = "Y" if option_text == q['correct'] else "N"
            row = {
                'Question ID (*required)': q['q_id'],
                'Locale ID (*required)': "English",
                'Domain ID (*required)': domain_id,
                'Active': 'Y',
                'Point Value': 1,
                'Variant Number (*required)': 1,
                'Question Type (*required)': "Multiple Choice",
                'Question Text (*required)': q['text'],
                'Answer Choice Number (*required)': opt_idx + 1,
                'Answer Choice Value': option_text,
                'Is Correct (*required)': is_correct,
            }
            rows.append(row)
            
    df_new = pd.DataFrame(rows)
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_new.to_excel(writer, sheet_name='Question', index=False)
        pd.DataFrame().to_excel(writer, sheet_name='Feedback', index=False)
        pd.DataFrame().to_excel(writer, sheet_name='Objectives', index=False)
    
    output.seek(0)
    return output

# --- Streamlit UI ---
st.set_page_config(page_title="LMS Question Converter", layout="centered")

st.title("📄 Word to SAP SF LMS Converter")
st.write("Upload your formatted Word document to generate the LMS import file. Safely handles variable option counts (2, 4, 6+), soft-returns, and image placeholders.")

exam_id = st.text_input("Enter Exam ID (e.g., MCT_2023)", "MCT_2023")
domain_id = st.text_input("Enter LMS Domain ID", "YOUR_DOMAIN")
uploaded_word = st.file_uploader("Upload Word Document (.docx)", type=["docx"])

if st.button("Convert Document"):
    if uploaded_word is not None and exam_id:
        with st.spinner("Dynamically parsing document..."):
            try:
                questions = parse_unlimited_options(uploaded_word, exam_id)
                
                if not questions:
                    st.error("No valid questions found. Please check the document format.")
                else:
                    excel_data = generate_lms_excel(questions, domain_id)
                    st.success(f"Successfully processed {len(questions)} questions!")
                    
                    st.download_button(
                        label="📥 Download LMS Excel File",
                        data=excel_data,
                        file_name=f"LMS_Import_{exam_id}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
            except Exception as e:
                st.error(f"An error occurred: {e}")
    else:
        st.warning("Please upload a Word document and provide an Exam ID.")
